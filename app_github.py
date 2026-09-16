"""
Inbound Tradebook 웹 대시보드
------------------------------------
OPEN(작업용) 시트를 표 형태로 그대로 옮겨서, 엑셀에서 행을 복사해
표에 바로 붙여넣거나(첫 셀 클릭 후 Ctrl+V) 셀을 직접 수정할 수 있습니다.

실행 방법:
    pip install streamlit pandas
    streamlit run app.py
"""

import sqlite3
import re
from io import BytesIO

import pandas as pd
import requests
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font

# ECOS(한국은행) API 인증키 — 발급받은 키를 여기 넣어두면 됩니다.
# 다른 사람과 파일을 공유하실 땐 이 값은 지우고 공유하세요.
ECOS_API_KEY = ""  # GitHub에 올릴 땐 비워두고, 로컬 실행 시 본인 키를 넣어 쓰세요

DB_PATH = "tradebook_v2.db"

# OPEN(작업용) 시트의 컬럼 순서 그대로 (엑셀에서 복사했을 때의 순서와 동일해야
# 표에 바로 붙여넣기가 됩니다)
BASE_COLUMNS = [
    "note",                    # 비고
    "kis_ref_no",               # KIS Ref. No.
    "next_valuation_date",      # Next Valuation Date
    "counterparty",              # Counterparty
    "trs_position",              # TRS Position
    "trade_date",                 # Trade Date
    "effective_date",             # Effective Date
    "final_valuation_date",       # Final Valuation Date
    "final_settlement_date",      # Final Settlement Date
    "underlying_code",            # Underlying (code)
    "underlying_description",     # Underlying (description)
    "number_of_units",            # Number of Units
    "initial_price",              # Initial Price
    "floating_rate_index",        # Floating Rate Index
    "spread",                     # spread
    "equity_notional_amount",     # Equity Notional Amount
    "net_number_of_units",        # Net Number of Units
]
RESET_COLUMNS = [f"reset_{i}" for i in range(1, 15)]  # Reset 1~14
EXCEL_COLUMNS = BASE_COLUMNS + RESET_COLUMNS  # 엑셀 붙여넣기 순서와 완전히 동일

# 엑셀에 없는, DB 관리용 컬럼 (표의 맨 뒤에 배치 — 붙여넣기 정렬을 깨지 않기 위함)
META_COLUMNS = ["book", "status", "id"]
EDITOR_COLUMNS = EXCEL_COLUMNS + META_COLUMNS

# UNWIND(작업용) — 원본 시트에서 연속된 4개 컬럼만 사용 (Sales Tax, Comms, FX Rate 제외)
UNWIND_EXCEL_COLUMNS = ["trs_code", "unwind_date", "unwind_size", "unwind_price"]
# 기본값이 있지만 예외적으로 수동 조정이 필요한 값들
UNWIND_OVERRIDE_COLUMNS = ["sales_tax", "settlement_days", "interest_rate_override"]
UNWIND_EDITOR_COLUMNS = UNWIND_EXCEL_COLUMNS + UNWIND_OVERRIDE_COLUMNS + ["id"]

# WORKDAY(휴일 캘린더) 시트 — A열=날짜, B열=휴일명
HOLIDAY_EXCEL_COLUMNS = ["date", "name"]
HOLIDAY_EDITOR_COLUMNS = HOLIDAY_EXCEL_COLUMNS + ["id"]


# ---------- DB ----------

def get_conn():
    conn = sqlite3.connect(DB_PATH)
    return conn


def init_db():
    conn = get_conn()
    reset_cols_sql = ", ".join(f"{c} TEXT" for c in RESET_COLUMNS)
    conn.executescript(
        f"""
        CREATE TABLE IF NOT EXISTS trades (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            kis_ref_no TEXT UNIQUE,
            book TEXT NOT NULL DEFAULT 'End Client',
            status TEXT NOT NULL DEFAULT 'open',
            note TEXT,
            next_valuation_date TEXT,
            counterparty TEXT,
            trs_position TEXT,
            trade_date TEXT,
            effective_date TEXT,
            final_valuation_date TEXT,
            final_settlement_date TEXT,
            underlying_code TEXT,
            underlying_description TEXT,
            number_of_units REAL,
            initial_price REAL,
            floating_rate_index TEXT,
            spread REAL,
            equity_notional_amount REAL,
            net_number_of_units REAL,
            {reset_cols_sql},
            partner_trade_id INTEGER,
            rollover_from_id INTEGER,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        );

        CREATE TABLE IF NOT EXISTS holidays (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            date TEXT UNIQUE,
            name TEXT
        );

        CREATE TABLE IF NOT EXISTS unwinds (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            trs_code TEXT,
            unwind_date TEXT,
            unwind_size REAL,
            unwind_price REAL,
            sales_tax TEXT NOT NULL DEFAULT 'N',
            settlement_days INTEGER NOT NULL DEFAULT 4,
            interest_rate_override REAL,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        );
        """
    )
    conn.commit()

    # ── 마이그레이션: 예전 버전 DB에 unwinds의 오버라이드 컬럼, holidays의 id 컬럼이
    #    없을 수 있음 (CREATE TABLE IF NOT EXISTS는 기존 테이블을 건드리지 않으므로) ──
    cur = conn.cursor()

    cur.execute("PRAGMA table_info(unwinds)")
    unwind_cols = {row[1] for row in cur.fetchall()}
    for col, ddl in [
        ("sales_tax", "ALTER TABLE unwinds ADD COLUMN sales_tax TEXT NOT NULL DEFAULT 'N'"),
        ("settlement_days", "ALTER TABLE unwinds ADD COLUMN settlement_days INTEGER NOT NULL DEFAULT 4"),
        ("interest_rate_override", "ALTER TABLE unwinds ADD COLUMN interest_rate_override REAL"),
    ]:
        if col not in unwind_cols:
            cur.execute(ddl)

    cur.execute("PRAGMA table_info(holidays)")
    holiday_cols = {row[1] for row in cur.fetchall()}
    if "id" not in holiday_cols:
        cur.executescript(
            """
            ALTER TABLE holidays RENAME TO holidays_old;
            CREATE TABLE holidays (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                date TEXT UNIQUE,
                name TEXT
            );
            INSERT INTO holidays (date, name) SELECT date, name FROM holidays_old;
            DROP TABLE holidays_old;
            """
        )

    conn.commit()
    conn.close()


def excel_serial_to_date(n) -> str:
    """엑셀 날짜 시리얼값(예: 46325)을 'YYYY-MM-DD' 문자열로 변환.
    엑셀은 1899-12-30을 0일로 취급(1900 윤년 버그 포함해서)하므로 그 기준으로 계산."""
    base = pd.Timestamp("1899-12-30")
    return (base + pd.Timedelta(days=float(n))).date().isoformat()


def clean_value(v, is_date: bool = False):
    """빈 문자열, '-', NaN 등을 None으로 통일. is_date=True면 엑셀 시리얼 숫자
    (예: 46325)나 다양한 날짜 문자열을 'YYYY-MM-DD'로 표준화."""
    if v is None:
        return None
    if isinstance(v, float) and pd.isna(v):
        return None
    s = str(v).strip()
    if s in ("", "-", "nan", "NaT", "None"):
        return None

    if is_date:
        # 순수 숫자(정수)면 엑셀 날짜 시리얼값으로 간주 — 그럴듯한 범위(1900~2200년대)만 변환
        try:
            f = float(s)
            if f == int(f) and 1 <= f <= 146000:
                return excel_serial_to_date(f)
        except ValueError:
            pass
        # 그 외 날짜 형식 문자열은 표준 형식으로 통일
        parsed = pd.to_datetime(s, errors="coerce")
        if pd.notna(parsed):
            return parsed.date().isoformat()

    return s


# 날짜로 취급해 자동 변환할 컬럼들
TRADE_DATE_COLUMNS = {
    "next_valuation_date", "trade_date", "effective_date",
    "final_valuation_date", "final_settlement_date",
} | set(RESET_COLUMNS)
UNWIND_DATE_COLUMNS = {"unwind_date"}
HOLIDAY_DATE_COLUMNS = {"date"}


def load_editor_df(book: str = None, counterparty: str = None, on_date=None) -> pd.DataFrame:
    conn = get_conn()
    query = f"SELECT {', '.join(EDITOR_COLUMNS)} FROM trades WHERE 1=1"
    params = []
    if book and book != "전체":
        query += " AND book = ?"
        params.append(book)
    if counterparty and counterparty != "전체":
        query += " AND counterparty = ?"
        params.append(counterparty)
    if on_date is not None:
        query += " AND date(created_at) = ?"
        params.append(str(on_date))
    query += " ORDER BY id"
    df = pd.read_sql(query, conn, params=params)
    conn.close()
    return df


def load_distinct_counterparties(book: str = None) -> list:
    conn = get_conn()
    query = "SELECT DISTINCT counterparty FROM trades WHERE counterparty IS NOT NULL"
    params = []
    if book and book != "전체":
        query += " AND book = ?"
        params.append(book)
    query += " ORDER BY counterparty"
    df = pd.read_sql(query, conn, params=params)
    conn.close()
    return df["counterparty"].dropna().tolist()


def apply_editor_changes(original_df: pd.DataFrame, editor_state: dict, default_book: str = "End Client"):
    """data_editor의 session_state(added/edited/deleted)를 DB에 반영."""
    conn = get_conn()
    cur = conn.cursor()
    n_added = n_edited = n_deleted = 0

    # 삭제된 행
    for row_idx in editor_state.get("deleted_rows", []):
        row_id = original_df.iloc[row_idx]["id"]
        if pd.notna(row_id):
            cur.execute("DELETE FROM trades WHERE id = ?", (int(row_id),))
            n_deleted += 1

    # 수정된 셀
    for row_idx, changes in editor_state.get("edited_rows", {}).items():
        row_id = original_df.iloc[int(row_idx)]["id"]
        if pd.isna(row_id) or not changes:
            continue
        clean_changes = {
            k: clean_value(v, is_date=k in TRADE_DATE_COLUMNS)
            for k, v in changes.items() if k != "id"
        }
        if not clean_changes:
            continue
        set_clause = ", ".join(f"{col} = ?" for col in clean_changes.keys())
        cur.execute(
            f"UPDATE trades SET {set_clause} WHERE id = ?",
            [*clean_changes.values(), int(row_id)],
        )
        n_edited += 1

    # 붙여넣기 또는 + 버튼으로 추가된 새 행
    for new_row in editor_state.get("added_rows", []):
        row = {
            k: clean_value(v, is_date=k in TRADE_DATE_COLUMNS)
            for k, v in new_row.items() if k != "id"
        }
        cols = [c for c in EDITOR_COLUMNS if c in row and c != "id"]
        if not cols:
            continue  # 완전히 빈 행
        if "book" not in row or row.get("book") is None:
            row["book"] = default_book
            if "book" not in cols:
                cols.append("book")

        # 같은 KIS Ref. No.가 이미 있으면 새로 추가하지 않고 업데이트
        # (재붙여넣기로 오타를 고치는 경우 UNIQUE 제약으로 앱이 죽는 것을 방지)
        existing_id = None
        if row.get("kis_ref_no"):
            cur.execute(
                "SELECT id FROM trades WHERE kis_ref_no = ?", (row["kis_ref_no"],)
            )
            found = cur.fetchone()
            if found:
                existing_id = found[0]

        if existing_id:
            set_clause = ", ".join(f"{c} = ?" for c in cols)
            cur.execute(
                f"UPDATE trades SET {set_clause} WHERE id = ?",
                [row.get(c) for c in cols] + [existing_id],
            )
            n_edited += 1
        else:
            placeholders = ", ".join(["?"] * len(cols))
            cur.execute(
                f"INSERT INTO trades ({', '.join(cols)}) VALUES ({placeholders})",
                [row.get(c) for c in cols],
            )
            n_added += 1

    conn.commit()
    conn.close()
    return n_added, n_edited, n_deleted


def load_trades(counterparty=None, book=None, status=None) -> pd.DataFrame:
    conn = get_conn()
    query = "SELECT * FROM trades WHERE 1=1"
    params = []
    if counterparty and counterparty != "전체":
        query += " AND counterparty LIKE ?"
        params.append(f"%{counterparty}%")
    if book and book != "전체":
        query += " AND book = ?"
        params.append(book)
    if status and status != "전체":
        query += " AND status = ?"
        params.append(status)
    query += " ORDER BY created_at DESC"
    df = pd.read_sql(query, conn, params=params)
    conn.close()
    return df


def load_unwound_totals() -> dict:
    """TRS코드(kis_ref_no)별 누적 청산 수량 합계 — OPEN(최종본) AV열과 동일한 로직."""
    conn = get_conn()
    try:
        df = pd.read_sql(
            "SELECT trs_code, SUM(unwind_size) AS total FROM unwinds "
            "WHERE trs_code IS NOT NULL GROUP BY trs_code",
            conn,
        )
    finally:
        conn.close()
    return dict(zip(df["trs_code"], df["total"]))


def load_upcoming_resets(days_ahead: int = 14, book: str = None, counterparty: str = None) -> pd.DataFrame:
    """모든 거래의 Reset 1~14 날짜를 펼쳐서, 오늘부터 days_ahead일 이내(지난 것 포함 -7일까지)
    예정된 리셋만 뽑아 정리."""
    trades_df = load_trades(counterparty=counterparty, book=book)
    if trades_df.empty:
        return pd.DataFrame()

    today = pd.Timestamp.now().normalize()
    window_end = today + pd.Timedelta(days=days_ahead)
    window_start = today - pd.Timedelta(days=7)  # 최근 놓친 것도 같이 보이게

    rows = []
    for _, trade in trades_df.iterrows():
        for i in range(1, 15):
            reset_date = pd.to_datetime(trade.get(f"reset_{i}"), errors="coerce")
            if pd.isna(reset_date):
                continue
            if window_start <= reset_date <= window_end:
                rows.append(
                    {
                        "Reset 예정일": reset_date.date(),
                        "D-day": (reset_date - today).days,
                        "TRS Code": trade.get("kis_ref_no"),
                        "Counterparty": trade.get("counterparty"),
                        "Underlying": trade.get("underlying_description"),
                        "Position": trade.get("trs_position"),
                        "Reset 회차": i,
                        "Book": trade.get("book"),
                        "상태": (
                            "⚠️ 지남" if reset_date < today
                            else ("🔴 오늘" if reset_date == today else "예정")
                        ),
                    }
                )

    if not rows:
        return pd.DataFrame()
    return pd.DataFrame(rows).sort_values("Reset 예정일")


# ---------- 휴일 캘린더 (WORKDAY 대체) ----------

def load_holiday_set() -> set:
    conn = get_conn()
    try:
        df = pd.read_sql("SELECT date FROM holidays", conn)
    finally:
        conn.close()
    out = set()
    for v in df["date"]:
        d = pd.to_datetime(v, errors="coerce")
        if pd.notna(d):
            out.add(d.date())
    return out


def load_holiday_editor_df() -> pd.DataFrame:
    conn = get_conn()
    df = pd.read_sql(
        f"SELECT {', '.join(HOLIDAY_EDITOR_COLUMNS)} FROM holidays ORDER BY date", conn
    )
    conn.close()
    return df


def apply_holiday_editor_changes(original_df: pd.DataFrame, editor_state: dict):
    conn = get_conn()
    cur = conn.cursor()
    n_added = n_edited = n_deleted = 0

    for row_idx in editor_state.get("deleted_rows", []):
        row_id = original_df.iloc[row_idx]["id"]
        if pd.notna(row_id):
            cur.execute("DELETE FROM holidays WHERE id = ?", (int(row_id),))
            n_deleted += 1

    for row_idx, changes in editor_state.get("edited_rows", {}).items():
        row_id = original_df.iloc[int(row_idx)]["id"]
        if pd.isna(row_id) or not changes:
            continue
        clean_changes = {
            k: clean_value(v, is_date=k in HOLIDAY_DATE_COLUMNS)
            for k, v in changes.items() if k != "id"
        }
        if not clean_changes:
            continue
        set_clause = ", ".join(f"{col} = ?" for col in clean_changes.keys())
        cur.execute(
            f"UPDATE holidays SET {set_clause} WHERE id = ?",
            [*clean_changes.values(), int(row_id)],
        )
        n_edited += 1

    for new_row in editor_state.get("added_rows", []):
        row = {
            k: clean_value(v, is_date=k in HOLIDAY_DATE_COLUMNS)
            for k, v in new_row.items() if k != "id"
        }
        cols = [c for c in HOLIDAY_EXCEL_COLUMNS if c in row]
        if not cols or "date" not in row or row.get("date") is None:
            continue  # 날짜 없는 빈 행은 스킵
        # 같은 날짜가 이미 있으면 업데이트 (중복 붙여넣기 방지)
        cur.execute("SELECT id FROM holidays WHERE date = ?", (row["date"],))
        found = cur.fetchone()
        if found:
            set_clause = ", ".join(f"{c} = ?" for c in cols)
            cur.execute(
                f"UPDATE holidays SET {set_clause} WHERE id = ?",
                [row.get(c) for c in cols] + [found[0]],
            )
            n_edited += 1
        else:
            placeholders = ", ".join(["?"] * len(cols))
            cur.execute(
                f"INSERT INTO holidays ({', '.join(cols)}) VALUES ({placeholders})",
                [row.get(c) for c in cols],
            )
            n_added += 1

    conn.commit()
    conn.close()
    return n_added, n_edited, n_deleted


def add_workdays(start_date, n_days: int, holidays: set):
    """엑셀 WORKDAY(start, n, holidays)와 동일한 로직: 토/일과 holidays를 건너뛰고 n 영업일 뒤 날짜."""
    if pd.isna(start_date) or n_days is None:
        return None
    d = pd.to_datetime(start_date, errors="coerce")
    if pd.isna(d):
        return None
    d = d.date()
    step = 1 if n_days >= 0 else -1
    remaining = abs(int(n_days))
    while remaining > 0:
        d = (pd.Timestamp(d) + pd.Timedelta(days=step)).date()
        if d.weekday() < 5 and d not in holidays:
            remaining -= 1
    return d


# ---------- UNWIND(최종본) 계산 로직 재현 ----------
# OPEN(최종본)/UNWIND(최종본) 수식을 그대로 옮긴 것. 참고:
#   Reset Price(Q), start accrual(U)는 원래 ^RESET(작업용) 시트에서 최근 리셋값을 찾고,
#   없으면 Initial Price / Effective Date로 대체(fallback)하는 로직인데,
#   RESET(작업용)을 아직 웹으로 안 옮겼으므로 지금은 항상 fallback 값을 씁니다.
#   Interest Rate(Y) = Benchmark(블룸버그 실시간) + Spread인데 블룸버그 연동 전이라
#   Benchmark는 사용자가 interest_rate_override에 직접 입력해야 계산됩니다.

def compute_unwind_result(trade: dict, unwind: dict, holidays: set) -> dict:
    fx_rate = 1.0  # OPEN(최종본)과 동일하게 항상 1로 고정

    unwind_size = float(unwind.get("unwind_size") or 0)
    unwind_price = float(unwind.get("unwind_price") or 0)
    comm = 0.0  # Comms는 항상 0이라는 전제로 컬럼 자체를 뺐음
    sales_tax = (unwind.get("sales_tax") or "N").upper()
    is_long = (trade.get("trs_position") or "").strip().lower() == "long"

    executed_notional = unwind_size * unwind_price  # F

    if sales_tax == "N":
        unwinding_notional = (
            executed_notional - executed_notional * comm
            if is_long
            else executed_notional + executed_notional * comm
        )
    else:
        unwinding_notional = (
            executed_notional - executed_notional * (comm + 0.0023)
            if is_long
            else executed_notional + executed_notional * comm
        )  # M

    settlement_days = int(unwind.get("settlement_days") or 4)
    settlement_date = add_workdays(unwind.get("unwind_date"), settlement_days, holidays)  # G

    number_of_units = float(trade.get("number_of_units") or 0)
    already_unwound = float(trade.get("_prior_unwound", 0))  # 이 거래의 기존 청산 수량 합
    balance = number_of_units - already_unwound  # H (Net No of Units)

    # Reset Price(Q) / start accrual(U): RESET(작업용) 연동 전이라 fallback 값 사용
    reset_price = float(trade.get("initial_price") or 0)  # Q fallback = Initial Price
    start_accrual = pd.to_datetime(trade.get("effective_date"), errors="coerce")  # U fallback

    # Interest Fixing Date(AQ) = start accrual의 2영업일 전 — 블룸버그 수식 내보내기에 사용
    fixing_date = None
    if pd.notna(start_accrual):
        fixing_date = add_workdays(start_accrual, -2, holidays)

    net_closing_price = (
        unwinding_notional / unwind_size / fx_rate if unwind_size else None
    )  # S

    if net_closing_price is None:
        equity_reset_payment = None
    else:
        equity_reset_payment = (
            -(reset_price - net_closing_price) * unwind_size
            if is_long
            else (reset_price - net_closing_price) * unwind_size
        )  # T

    end_accrual = pd.to_datetime(settlement_date, errors="coerce")  # V
    unwind_notional = reset_price * unwind_size  # W

    no_of_days = None
    if pd.notna(start_accrual) and pd.notna(end_accrual):
        no_of_days = (end_accrual - start_accrual).days  # X

    interest_rate = unwind.get("interest_rate_override")  # Y — 수동 입력 없으면 계산 불가
    interest_rate = float(interest_rate) if interest_rate not in (None, "") else None

    floating_rate_payments = None
    if interest_rate is not None and no_of_days is not None:
        day_count = 365 if fx_rate == 1 else 360
        floating_rate_payments = (
            -unwind_notional * interest_rate * no_of_days / day_count
            if is_long
            else unwind_notional * interest_rate * no_of_days / day_count
        )  # Z

    unwind_payment = None
    if equity_reset_payment is not None and floating_rate_payments is not None:
        unwind_payment = equity_reset_payment + floating_rate_payments  # O
    elif equity_reset_payment is not None:
        unwind_payment = equity_reset_payment  # 이자율 미입력 시 이자분 제외한 잠정값

    unwind_payment_underlying_ccy = (
        unwind_payment * fx_rate if unwind_payment is not None else None
    )  # P

    direction = None
    if unwind_payment is not None:
        direction = "KIS to Pay" if unwind_payment > 0 else "KIS to Rec"  # N

    return {
        # UNWIND_Raw / UNWIND(최종본) A~Z 컬럼 그대로 (엑셀에 바로 붙여넣을 수 있는 순서)
        "KIS Ref. No.": unwind.get("trs_code"),                      # A
        "Unwind Valuation Date": unwind.get("unwind_date"),          # B
        "Unwinding Size": unwind_size,                                # C
        "FX Rate": fx_rate,                                           # D
        "Comm": comm,                                                 # E
        "Executed Notional": round(executed_notional, 2),            # F
        "Settlement Date": settlement_date,                           # G
        "Balance": round(balance, 2),                                 # H
        "Counterparty": trade.get("counterparty"),                    # I
        "Swap Position": trade.get("trs_position"),                   # J
        "Underlying (description)": trade.get("underlying_description"),  # K
        "Underlying (code)": trade.get("underlying_code"),            # L
        "Unwinding Notional": round(unwinding_notional, 2),           # M
        "Direction": direction,                                       # N
        "Unwind Payment (Swap CCY)": round(unwind_payment, 2) if unwind_payment is not None else None,  # O
        "Unwind Payment (Underlying CCY)": round(unwind_payment_underlying_ccy, 2) if unwind_payment_underlying_ccy is not None else None,  # P
        "Reset Price": round(reset_price, 4),                         # Q
        "Executed Price": unwind_price,                               # R
        "Net Closing Price": round(net_closing_price, 4) if net_closing_price is not None else None,  # S
        "Equity Reset Payment": round(equity_reset_payment, 2) if equity_reset_payment is not None else None,  # T
        "start accrual": start_accrual.date() if pd.notna(start_accrual) else None,  # U
        "end accrual": end_accrual.date() if pd.notna(end_accrual) else None,  # V
        "Unwind Notional": round(unwind_notional, 2),                 # W
        "No. of Days": no_of_days,                                    # X
        "Interest Rate": interest_rate,                               # Y
        "Floating Rate Payments": round(floating_rate_payments, 2) if floating_rate_payments is not None else None,  # Z
        # 참고용 (Unwind_Raw엔 없지만 화면 확인용으로 같이 보여줌)
        "Floating Rate Index": trade.get("floating_rate_index"),
        "Interest Fixing Date": fixing_date,
        "Spread": trade.get("spread"),
    }


UNWIND_RAW_COLUMNS = [
    "KIS Ref. No.", "Unwind Valuation Date", "Unwinding Size", "FX Rate", "Comm",
    "Executed Notional", "Settlement Date", "Balance", "Counterparty", "Swap Position",
    "Underlying (description)", "Underlying (code)", "Unwinding Notional", "Direction",
    "Unwind Payment (Swap CCY)", "Unwind Payment (Underlying CCY)", "Reset Price",
    "Executed Price", "Net Closing Price", "Equity Reset Payment", "start accrual",
    "end accrual", "Unwind Notional", "No. of Days", "Interest Rate",
    "Floating Rate Payments",
]


# ---------- 참고용 시세/금리 조회 (블룸버그 아님 — 확인용) ----------

def bloomberg_code_to_yahoo_ticker(code: str):
    """'688187 CH Equity' 같은 블룸버그 티커를 야후 파이낸스 형식으로 변환.
    6로 시작하면 상해(.SS), 0/3으로 시작하면 심천(.SZ)."""
    if not code:
        return None
    digits = code.strip().split()[0]
    if not digits.isdigit():
        return None
    prefix = digits[0]
    if prefix == "6":
        return f"{digits}.SS"
    if prefix in ("0", "3"):
        return f"{digits}.SZ"
    return None


def fetch_reference_last_price(bloomberg_code: str):
    """yfinance로 참고용 종가 조회. 반환: (가격, 날짜) 또는 (None, 에러메시지)."""
    try:
        import yfinance as yf
    except ImportError:
        return None, "yfinance가 설치되어 있지 않습니다 (pip install yfinance)"

    yahoo_ticker = bloomberg_code_to_yahoo_ticker(bloomberg_code)
    if not yahoo_ticker:
        return None, f"'{bloomberg_code}'를 야후 티커로 변환하지 못했습니다"

    try:
        hist = yf.Ticker(yahoo_ticker).history(period="5d")
        if hist.empty:
            return None, f"{yahoo_ticker}: 데이터 없음"
        last_close = float(hist["Close"].iloc[-1])
        last_date = str(hist.index[-1].date())
        return (last_close, last_date, yahoo_ticker), None
    except Exception as e:
        return None, str(e)


def fetch_reference_cd91_rate(api_key: str, date_str: str):
    """한국은행 ECOS에서 CD91(91일물) 금리 참고 조회. date_str: 'YYYYMMDD'.
    반환: (금리, None) 또는 (None, 에러메시지)."""
    import requests

    url = (
        f"https://ecos.bok.or.kr/api/StatisticSearch/{api_key}/json/kr/1/5/"
        f"817Y002/D/{date_str}/{date_str}/010502000"
    )
    try:
        res = requests.get(url, timeout=10).json()
        rows = res.get("StatisticSearch", {}).get("row", [])
        if not rows:
            return None, f"{date_str}: ECOS에 데이터 없음 (휴장일이거나 항목코드 확인 필요)"
        return float(rows[0]["DATA_VALUE"]), None
    except Exception as e:
        return None, str(e)


def build_unwind_raw_workbook(results: list) -> bytes:
    """Unwind_Raw 시트(A~Z, 7행부터 데이터)에 그대로 붙여넣을 수 있는 형태로 내보내기.
    헤더도 6행에 원본과 동일하게 넣어서, 통째로 복사해 옮기기 쉽게 함."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Unwind_Raw용"
    font = Font(name="Arial", size=10)
    header_font = Font(name="Arial", size=10, bold=True)

    for c, col in enumerate(UNWIND_RAW_COLUMNS, start=1):
        cell = ws.cell(row=6, column=c, value=col)
        cell.font = header_font

    for r, row in enumerate(results, start=7):
        for c, col in enumerate(UNWIND_RAW_COLUMNS, start=1):
            v = row.get(col)
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = font
            if col in ("Unwind Valuation Date", "Settlement Date", "start accrual", "end accrual") and v is not None:
                cell.number_format = "yyyy-mm-dd"

    for i in range(1, len(UNWIND_RAW_COLUMNS) + 1):
        ws.column_dimensions[chr(64 + i) if i <= 26 else "A"].width = 16

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_bloomberg_formula_workbook(rows: list) -> bytes:
    """Interest Rate(Net Financing Rate)를 실제 블룸버그 수식(_xll.BDH)으로 계산하는
    엑셀 파일 생성. 블룸버그 애드인이 깔린 PC(터미널 PC)에서 열어야 값이 채워짐 —
    이 샌드박스에는 애드인이 없어 여기선 수식 문자열만 정확히 써넣고 계산은 하지 않음."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Interest Rate 조회"

    headers = [
        "TRS Code", "Floating Rate Index", "Interest Fixing Date",
        "Spread", "Net Financing Rate (Benchmark+Spread)",
    ]
    font = Font(name="Arial", size=10)
    header_font = Font(name="Arial", size=10, bold=True)
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = header_font

    for r, row in enumerate(rows, start=2):
        ws.cell(row=r, column=1, value=row.get("trs_code")).font = font
        ws.cell(row=r, column=2, value=row.get("floating_rate_index")).font = font

        fixing_date = row.get("fixing_date")
        date_cell = ws.cell(row=r, column=3)
        if fixing_date is not None:
            date_cell.value = pd.Timestamp(fixing_date).to_pydatetime()
            date_cell.number_format = "yyyy-mm-dd"
        date_cell.font = font

        spread_cell = ws.cell(row=r, column=4, value=row.get("spread"))
        spread_cell.font = font

        # 원본 OPEN(최종본) AL/AR 수식 그대로: BDH(지수,"PX_LAST",고정일,고정일,"DAYS=A")/100 + Spread
        formula = (
            f'=IF(B{r}="","",_xll.BDH(B{r},"PX_LAST",C{r},C{r},"DAYS=A")/100+D{r})'
        )
        formula_cell = ws.cell(row=r, column=5, value=formula)
        formula_cell.font = font

    for col_letter, width in zip("ABCDE", [16, 20, 18, 10, 30]):
        ws.column_dimensions[col_letter].width = width

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- UNWIND ----------

def load_unwind_editor_df(on_date=None, search_text: str = None) -> pd.DataFrame:
    conn = get_conn()
    cols = UNWIND_EDITOR_COLUMNS + ["created_at"] if search_text else UNWIND_EDITOR_COLUMNS
    query = f"SELECT {', '.join('u.' + c for c in cols)} FROM unwinds u WHERE 1=1"
    params = []
    if search_text:
        query += """ AND (
            u.trs_code LIKE ?
            OR u.trs_code IN (
                SELECT kis_ref_no FROM trades
                WHERE underlying_code LIKE ?
                   OR underlying_description LIKE ?
                   OR counterparty LIKE ?
            )
        )"""
        like = f"%{search_text}%"
        params.extend([like, like, like, like])
    elif on_date is not None:
        query += " AND date(u.created_at) = ?"
        params.append(str(on_date))
    query += " ORDER BY u.id"
    df = pd.read_sql(query, conn, params=params)
    conn.close()
    return df


def apply_unwind_editor_changes(original_df: pd.DataFrame, editor_state: dict):
    conn = get_conn()
    cur = conn.cursor()
    n_added = n_edited = n_deleted = 0

    for row_idx in editor_state.get("deleted_rows", []):
        row_id = original_df.iloc[row_idx]["id"]
        if pd.notna(row_id):
            cur.execute("DELETE FROM unwinds WHERE id = ?", (int(row_id),))
            n_deleted += 1

    for row_idx, changes in editor_state.get("edited_rows", {}).items():
        row_id = original_df.iloc[int(row_idx)]["id"]
        if pd.isna(row_id) or not changes:
            continue
        clean_changes = {
            k: clean_value(v, is_date=k in UNWIND_DATE_COLUMNS)
            for k, v in changes.items() if k != "id"
        }
        if not clean_changes:
            continue
        set_clause = ", ".join(f"{col} = ?" for col in clean_changes.keys())
        cur.execute(
            f"UPDATE unwinds SET {set_clause} WHERE id = ?",
            [*clean_changes.values(), int(row_id)],
        )
        n_edited += 1

    for new_row in editor_state.get("added_rows", []):
        row = {
            k: clean_value(v, is_date=k in UNWIND_DATE_COLUMNS)
            for k, v in new_row.items() if k != "id"
        }
        cols = [c for c in UNWIND_EXCEL_COLUMNS + UNWIND_OVERRIDE_COLUMNS if c in row]
        if not cols:
            continue
        placeholders = ", ".join(["?"] * len(cols))
        cur.execute(
            f"INSERT INTO unwinds ({', '.join(cols)}) VALUES ({placeholders})",
            [row.get(c) for c in cols],
        )
        n_added += 1

    conn.commit()
    conn.close()
    return n_added, n_edited, n_deleted


# ---------- yfinance / ECOS 연동 (Reset 준비용) ----------

def bloomberg_to_yfinance_ticker(underlying_code: str):
    """'688187 CH Equity' -> '688187.SS' (6으로 시작=상하이) / '300124 CH Equity' -> '300124.SZ' (0,3=선전).
    중국 A주가 아니면 None 반환."""
    m = re.match(r"(\d{6})\s*CH\s*Equity", (underlying_code or "").strip(), re.IGNORECASE)
    if not m:
        return None
    code = m.group(1)
    if code.startswith("6"):
        return f"{code}.SS"
    if code.startswith(("0", "3")):
        return f"{code}.SZ"
    return None


def fetch_yf_price_asof(ticker: str, target_date, verify_ssl: bool = True):
    """야후 파이낸스에서 target_date 시점(그 날짜가 휴일/주말이면 그 이전 가장 최근
    거래일) 종가를 조회. 엑셀 BDH의 DAYS=A와 같은 동작.
    (price, actual_date, company_name, error) 반환."""
    headers = {"User-Agent": "Mozilla/5.0"}
    target_ts = pd.Timestamp(target_date)

    # 공휴일 연휴 대비 여유있게 앞뒤로 넉넉히 잡음 (설/추석 등 최대 5일 연휴 고려해 14일 버퍼)
    period1 = int((target_ts - pd.Timedelta(days=14)).timestamp())
    period2 = int((target_ts + pd.Timedelta(days=1)).timestamp())

    try:
        chart_url = (
            f"https://query1.finance.yahoo.com/v8/finance/chart/{ticker}"
            f"?period1={period1}&period2={period2}&interval=1d"
        )
        resp = requests.get(chart_url, headers=headers, timeout=10, verify=verify_ssl)
        resp.raise_for_status()
        data = resp.json()

        result = data.get("chart", {}).get("result")
        if not result:
            err = data.get("chart", {}).get("error", {}).get("description", "데이터 없음")
            return None, None, None, err

        result = result[0]
        timestamps = result.get("timestamp", [])
        closes = result.get("indicators", {}).get("quote", [{}])[0].get("close", [])
        company_name = result.get("meta", {}).get("longName") or result.get("meta", {}).get("shortName")

        # target_date 이하(그 날짜 포함)인 것 중 가장 최근 거래일을 찾음
        candidates = [
            (ts, c) for ts, c in zip(timestamps, closes)
            if c is not None and pd.Timestamp(ts, unit="s").date() <= target_ts.date()
        ]
        if not candidates:
            return None, None, company_name, f"{target_ts.date()} 이전 14일 내 거래 데이터 없음"

        last_ts, last_close = max(candidates, key=lambda x: x[0])
        actual_date = pd.Timestamp(last_ts, unit="s").date().isoformat()
        return round(float(last_close), 4), actual_date, company_name, None
    except Exception as e:
        return None, None, None, str(e)



def ecos_list_items(stat_code: str, api_key: str = ECOS_API_KEY, verify_ssl: bool = True):
    """통계표(stat_code) 안의 전체 세부 항목 코드/이름 목록 조회 — 정확한 item code
    확인용. 예: ecos_list_items('817Y002')로 CD91일물 항목코드를 직접 확인."""
    url = (
        f"https://ecos.bok.or.kr/api/StatisticItemList/{api_key}/json/kr/1/100/{stat_code}"
    )
    resp = requests.get(url, timeout=10, verify=verify_ssl)
    resp.raise_for_status()
    data = resp.json()
    rows = data.get("StatisticItemList", {}).get("row", [])
    return pd.DataFrame(rows)


def fetch_ecos_series(
    stat_code: str, item_code: str, date_str: str,
    api_key: str = ECOS_API_KEY, verify_ssl: bool = True,
):
    """단일 일자 통계값 조회. date_str은 'YYYYMMDD'. (value, error) 반환."""
    url = (
        f"https://ecos.bok.or.kr/api/StatisticSearch/{api_key}/json/kr/1/5/"
        f"{stat_code}/D/{date_str}/{date_str}/{item_code}"
    )
    try:
        resp = requests.get(url, timeout=10, verify=verify_ssl)
        resp.raise_for_status()
        data = resp.json()
        if "RESULT" in data:
            return None, data["RESULT"].get("MESSAGE", "알 수 없는 오류")
        rows = data.get("StatisticSearch", {}).get("row", [])
        if not rows:
            return None, "해당 날짜 데이터 없음 (공휴일이었을 수 있음)"
        return float(rows[0]["DATA_VALUE"]), None
    except Exception as e:
        return None, str(e)


# ---------- UI ----------

def date_navigator(key: str):
    """TMS 스타일 날짜 선택기: ◀ [날짜] ▶ 오늘 버튼. 선택된 date 객체를 반환."""
    state_key = f"{key}_date"
    if state_key not in st.session_state:
        st.session_state[state_key] = pd.Timestamp.now().date()

    col1, col2, col3, col4 = st.columns([1, 3, 1, 1.2])
    with col1:
        if st.button("◀", key=f"{key}_prev"):
            st.session_state[state_key] -= pd.Timedelta(days=1)
    with col2:
        st.session_state[state_key] = st.date_input(
            "날짜", value=st.session_state[state_key], key=f"{key}_input",
            label_visibility="collapsed",
        )
    with col3:
        if st.button("▶", key=f"{key}_next"):
            st.session_state[state_key] += pd.Timedelta(days=1)
    with col4:
        if st.button("오늘", key=f"{key}_today"):
            st.session_state[state_key] = pd.Timestamp.now().date()

    return st.session_state[state_key]


def seed_if_empty():
    """배포 환경이 재시작되어 DB가 비어있을 때, 가상 샘플 데이터를 자동으로 채운다.
    (심사위원이 언제 접속하든 빈 화면이 아니라 풍부하게 동작하는 화면을 보도록 하기 위한 안전장치)
    실제 거래 데이터는 전혀 사용하지 않으며, 계좌번호·펀드명·티커·거래상대방 모두 가상값이다."""
    import datetime as _dt

    conn = get_conn()
    cur = conn.cursor()
    cur.execute("SELECT COUNT(*) FROM trades")
    if cur.fetchone()[0] > 0:
        conn.close()
        return

    today = _dt.date.today()

    def d(offset):
        return (today + _dt.timedelta(days=offset)).isoformat()

    cols = [
        "kis_ref_no", "book", "note", "counterparty", "trs_position", "trade_date",
        "effective_date", "underlying_code", "underlying_description", "number_of_units",
        "initial_price", "floating_rate_index", "spread", "equity_notional_amount",
        "net_number_of_units", "reset_1", "reset_2", "reset_3", "reset_4",
    ]
    placeholders = ",".join(["?"] * len(cols))

    etf_letters = list("ABCDEFGHIJ")  # ETF A ~ ETF J, 10종 (가상)
    btb_counterparties = ["GS", "Nomura", "JPM-Sample", "Citi-Sample"]
    # 리셋 패턴을 다양하게: 지난 리셋 포함/다가오는 리셋만/리셋 여러 회차 등 인덱스별로 변주
    reset_patterns = [
        [d(-30), d(3), None, None],
        [d(8), None, None, None],
        [d(-45), d(-2), d(13), None],
        [d(20), None, None, None],
        [d(-60), d(-15), d(30), None],
        [d(5), None, None, None],
        [d(-90), d(-30), d(0), d(30)],
        [d(11), None, None, None],
        [d(-20), d(9), None, None],
        [d(2), None, None, None],
    ]

    for i, letter in enumerate(etf_letters):
        units = 5000 + i * 1500
        price = 10000.0 + i * 1800.0
        notional = round(units * price)
        counterparty_btb = btb_counterparties[i % len(btb_counterparties)]
        trade_offset = -30 - i * 12
        resets = reset_patterns[i % len(reset_patterns)]

        end_row = (
            f"TRS26SP{i+1:02d}01", "End Client", f"ETF {letter} 신규", "A자산운용", "LONG",
            d(trade_offset), d(trade_offset + 1), f"{i+1:06d} XX Equity", f"샘플종목 {letter}",
            units, price, "KWCDC Curncy", -0.004 - (i % 3) * 0.001, notional, resets,
        )
        btb_row = (
            f"TRS26SP{i+1:02d}02", "BTB", f"ETF {letter} 짝꿍({counterparty_btb})", counterparty_btb,
            "LONG", d(trade_offset), d(trade_offset + 1), f"{i+1:06d} XX Equity", f"샘플종목 {letter}",
            units, price, "KWCDC Curncy", 0.008 + (i % 3) * 0.001, notional, resets,
        )

        for ref, book, note, cp, pos, tdate, edate, ucode, udesc, u, p, fri, spread, notion, rst in (end_row, btb_row):
            vals = [ref, book, note, cp, pos, tdate, edate, ucode, udesc, u, p, fri, spread, notion, u] + rst
            cur.execute(
                f"INSERT OR IGNORE INTO trades ({','.join(cols)}) VALUES ({placeholders})",
                vals,
            )

    # 휴일 6건 — 연중 분산
    holidays = [
        (d(5), "샘플 공휴일1"), (d(20), "샘플 공휴일2"), (d(45), "샘플 공휴일3"),
        (d(-10), "샘플 공휴일4(과거)"), (d(70), "샘플 공휴일5"), (d(95), "샘플 공휴일6"),
    ]
    for hd, name in holidays:
        cur.execute("INSERT OR IGNORE INTO holidays (date, name) VALUES (?,?)", (hd, name))

    # UNWIND 샘플 6건 — 일부 ETF는 부분 청산된 상태로 시연
    unwind_targets = [
        ("TRS26SP0501", 5000, 12500.0),  # ETF E End Client (인덱스 4)
        ("TRS26SP0502", 5000, 12500.0),  # ETF E BTB
        ("TRS26SP0701", 3000, 22000.0),  # ETF G End Client (인덱스 6)
        ("TRS26SP0702", 3000, 22000.0),  # ETF G BTB
        ("TRS26SP0901", 2000, 26200.0),  # ETF I End Client (인덱스 8)
        ("TRS26SP0902", 2000, 26200.0),  # ETF I BTB
    ]
    # 위 kis_ref_no 포맷을 실제 생성 규칙(TRS26SP{idx:02d}0{1|2})에 맞춰 보정
    unwind_targets = [
        (f"TRS26SP{5:02d}01", 5000, 12500.0),
        (f"TRS26SP{5:02d}02", 5000, 12500.0),
        (f"TRS26SP{7:02d}01", 3000, 22000.0),
        (f"TRS26SP{7:02d}02", 3000, 22000.0),
        (f"TRS26SP{9:02d}01", 2000, 26200.0),
        (f"TRS26SP{9:02d}02", 2000, 26200.0),
    ]
    for trs_code, u_size, u_price in unwind_targets:
        cur.execute(
            "INSERT INTO unwinds (trs_code, unwind_date, unwind_size, unwind_price, "
            "sales_tax, settlement_days, interest_rate_override) VALUES (?,?,?,?,?,?,?)",
            (trs_code, d(-3), u_size, u_price, "N", 4, None),
        )

    conn.commit()
    conn.close()

st.set_page_config(page_title="Inbound Tradebook", layout="wide")
init_db()
seed_if_empty()

st.title("📋 Inbound Tradebook")

st.sidebar.markdown("### Book")
book_filter = st.sidebar.radio("Book", ["End Client", "BTB"], label_visibility="collapsed")

counterparty_filter = "전체"
if book_filter == "BTB":
    cps = load_distinct_counterparties("BTB")
    counterparty_filter = st.sidebar.selectbox(
        "Counterparty (BTB 세부)", ["전체"] + cps
    )

st.sidebar.divider()

page = st.sidebar.radio(
    "메뉴",
    [
        "다가오는 Reset",
        "신규 거래 입력",
        "거래 조회/수정/삭제",
        "UNWIND 입력/수정",
        "휴일 관리(WORKDAY)",
        "블룸버그 대체 데이터 확인",
        "조회",
    ],
)

if page == "다가오는 Reset":
    st.subheader(f"🔔 다가오는 Reset 안내 · {book_filter}" + (f" / {counterparty_filter}" if counterparty_filter != "전체" else ""))

    days_ahead = st.slider("며칠 이내 예정된 Reset까지 볼까요?", 1, 60, 14)
    upcoming = load_upcoming_resets(days_ahead, book=book_filter, counterparty=counterparty_filter)

    if upcoming.empty:
        st.info(f"앞으로 {days_ahead}일 이내(및 최근 7일간 지난 것 포함) 예정된 Reset이 없습니다.")
    else:
        overdue = upcoming[upcoming["상태"] == "⚠️ 지남"]
        if not overdue.empty:
            st.warning(f"⚠️ 이미 지난 Reset이 {len(overdue)}건 있습니다 — 처리 여부 확인하세요.")

        st.dataframe(upcoming, use_container_width=True, hide_index=True)
        st.caption(f"총 {len(upcoming)}건 (거래별 Reset 1~14를 모두 펼쳐서 표시)")

elif page == "신규 거래 입력":
    st.subheader(f"➕ 신규 거래 입력 · {book_filter}" + (f" / {counterparty_filter}" if counterparty_filter != "전체" else ""))
    st.caption(
        "선택한 날짜에 입력된 신규 거래만 보여요. 엑셀에서 행을 복사한 뒤, 맨 왼쪽('비고') "
        "빈 셀을 클릭하고 Ctrl+V로 붙여넣으세요. 기존 거래를 고치거나 지우려면 "
        "'거래 조회/수정/삭제' 메뉴를 쓰세요."
    )

    new_entry_date = date_navigator("new_entry")
    day_df = load_editor_df(book=book_filter, counterparty=counterparty_filter, on_date=new_entry_date)

    editor_key = f"new_entry_editor_{new_entry_date}"
    new_edited_df = st.data_editor(
        day_df,
        num_rows="dynamic",
        use_container_width=True,
        hide_index=True,
        key=editor_key,
        column_config={
            "id": st.column_config.NumberColumn("ID", disabled=True),
            "book": st.column_config.SelectboxColumn("Book", options=["End Client", "BTB"]),
            "status": st.column_config.SelectboxColumn(
                "상태", options=["open", "reset", "unwind", "rollover"]
            ),
            "trs_position": st.column_config.TextColumn("TRS Position"),
        },
    )

    new_state = st.session_state.get(editor_key, {})
    n_add = len(new_state.get("added_rows", []))
    n_edit = len(new_state.get("edited_rows", {}))
    n_del = len(new_state.get("deleted_rows", []))
    st.caption(f"[추가] {n_add}  [변경] {n_edit}  [삭제] {n_del}  (새로 추가하는 행은 실제 저장 시각 기준으로 오늘 날짜에 들어갑니다)")

    if st.button("Save Changes", type="primary", key="save_new_entry"):
        n_added, n_edited, n_deleted = apply_editor_changes(
            day_df, new_state, default_book=book_filter
        )
        st.success(f"저장 완료 — 추가 {n_added} / 변경 {n_edited} / 삭제 {n_deleted}")
        st.rerun()

elif page == "거래 조회/수정/삭제":
    st.subheader(f"거래 조회/수정/삭제 · {book_filter}" + (f" / {counterparty_filter}" if counterparty_filter != "전체" else ""))
    st.caption(
        "이미 저장된 거래를 고치거나 지우는 화면이에요. 셀을 더블클릭해 수정하거나, 행을 "
        "선택해 삭제 아이콘으로 지울 수 있습니다. 새 거래는 '신규 거래 입력' 메뉴를 쓰세요."
    )

    editor_df = load_editor_df(book=book_filter, counterparty=counterparty_filter)

    edited_df = st.data_editor(
        editor_df,
        num_rows="dynamic",
        use_container_width=True,
        hide_index=True,
        key="trade_editor",
        column_config={
            "id": st.column_config.NumberColumn("ID", disabled=True),
            "book": st.column_config.SelectboxColumn("Book", options=["End Client", "BTB"]),
            "status": st.column_config.SelectboxColumn(
                "상태", options=["open", "reset", "unwind", "rollover"]
            ),
            "trs_position": st.column_config.TextColumn("TRS Position"),
        },
    )

    state = st.session_state.get("trade_editor", {})
    n_add = len(state.get("added_rows", []))
    n_edit = len(state.get("edited_rows", {}))
    n_del = len(state.get("deleted_rows", []))
    st.caption(f"[추가] {n_add}  [변경] {n_edit}  [삭제] {n_del}")

    if st.button("Save Changes", type="primary"):
        n_added, n_edited, n_deleted = apply_editor_changes(editor_df, state, default_book=book_filter)
        st.success(f"저장 완료 — 추가 {n_added} / 변경 {n_edited} / 삭제 {n_deleted}")
        st.rerun()

elif page == "UNWIND 입력/수정":
    st.subheader("UNWIND(작업용)")
    st.caption(
        "엑셀 UNWIND(작업용) 시트에서 TRS Code, Unwind Date, Unwind Size, Unwind Price "
        "4개 컬럼(연속된 범위)만 선택해서 복사한 뒤, 표의 'trs_code' 빈 셀에 붙여넣으세요."
    )

    search_trs = st.text_input(
        "🔍 TRS Code / 티커 / 거래처로 검색 (날짜 상관없이 전체에서 찾기 — 수정/삭제할 때 사용)"
    )

    if search_trs:
        st.caption(f"'{search_trs}' 검색 결과 — 날짜 무관 전체. 여기서 바로 수정/삭제 가능합니다.")
        unwind_df = load_unwind_editor_df(search_text=search_trs)
        unwind_editor_key = f"unwind_search_editor_{search_trs}"
    else:
        unwind_date_selected = date_navigator("unwind_entry")
        st.caption(f"{unwind_date_selected} 에 저장된 UNWIND 건만 보여요.")
        unwind_df = load_unwind_editor_df(on_date=unwind_date_selected)
        unwind_editor_key = f"unwind_editor_{unwind_date_selected}"

    edited_unwind_df = st.data_editor(
        unwind_df,
        num_rows="dynamic",
        use_container_width=True,
        hide_index=True,
        key=unwind_editor_key,
        column_config={
            "id": st.column_config.NumberColumn("ID", disabled=True),
            "created_at": st.column_config.TextColumn("입력일시", disabled=True),
            "sales_tax": st.column_config.SelectboxColumn(
                "Sales Tax (기본 N)", options=["N", "Y"]
            ),
            "settlement_days": st.column_config.NumberColumn(
                "Settlement 영업일 (기본 4)", min_value=0, max_value=10, step=1
            ),
            "interest_rate_override": st.column_config.NumberColumn(
                "Interest Rate 수동입력 (Benchmark+Spread, 블룸버그 미연동으로 직접 입력 필요)",
                format="%.4f",
            ),
        },
    )

    u_state = st.session_state.get(unwind_editor_key, {})
    u_add = len(u_state.get("added_rows", []))
    u_edit = len(u_state.get("edited_rows", {}))
    u_del = len(u_state.get("deleted_rows", []))
    st.caption(f"[추가] {u_add}  [변경] {u_edit}  [삭제] {u_del}  (새로 추가하는 행은 실제 저장 시각 기준 오늘 날짜로 들어갑니다)")

    if st.button("Save Changes", type="primary", key="save_unwind"):
        n_added, n_edited, n_deleted = apply_unwind_editor_changes(unwind_df, u_state)
        st.success(f"저장 완료 — 추가 {n_added} / 변경 {n_edited} / 삭제 {n_deleted}")
        st.rerun()

    st.divider()
    st.subheader("UNWIND 계산 결과 (UNWIND(최종본) 재현)")
    st.caption(
        "⚠️ Reset Price / start accrual은 RESET(작업용) 연동 전이라 각각 Initial Price / "
        "Effective Date로 대체 계산됩니다. Interest Rate는 위 표에 직접 입력해야 이자 관련 "
        "항목(Floating Rate Payments, Unwind Payment, Direction)까지 계산됩니다."
    )

    saved_unwinds = load_unwind_editor_df()
    trades_all = load_trades(counterparty=counterparty_filter, book=book_filter)

    if saved_unwinds.empty or trades_all.empty:
        st.info("계산할 UNWIND 데이터 또는 매칭되는 거래가 아직 없습니다.")
    else:
        holidays = load_holiday_set()
        trades_by_ref = {
            row["kis_ref_no"]: row.to_dict() for _, row in trades_all.iterrows()
        }
        # TRS코드별 누적 청산 수량(AV 로직: 전체 unwind 합계)
        total_unwound = (
            saved_unwinds.dropna(subset=["trs_code"])
            .groupby("trs_code")["unwind_size"]
            .sum()
        )

        results = []
        results_raw = []  # (trade dict, unwind dict) — 참고용 재계산에 사용
        for _, urow in saved_unwinds.iterrows():
            trs_code = urow.get("trs_code")
            trade = trades_by_ref.get(trs_code)
            if not trade:
                continue
            trade = dict(trade)
            trade["_prior_unwound"] = float(total_unwound.get(trs_code, 0))
            results.append(compute_unwind_result(trade, urow.to_dict(), holidays))
            results_raw.append((trade, urow.to_dict()))

        if results:
            results_df = pd.DataFrame(results)
            st.dataframe(
                results_df[UNWIND_RAW_COLUMNS], use_container_width=True, hide_index=True
            )
            st.caption("UNWIND(최종본) A~Z와 동일한 컬럼 순서예요.")

            with st.expander("📡 참고용 Last Price / CD91 금리 조회 (블룸버그 아님, 확인용)"):
                st.caption(
                    "yfinance(Last Price)와 한국은행 ECOS(CD91 금리)로 참고용 값을 가져와요. "
                    "실제 정산에 쓰는 공식 값은 블룸버그 기준이니, 여기 값은 크로스체크 용도로만 "
                    "써주세요."
                )
                ref_options = list(range(len(results)))
                ref_idx = st.selectbox(
                    "조회할 건",
                    ref_options,
                    format_func=lambda i: f"{results[i]['KIS Ref. No.']} · {results[i]['Unwind Valuation Date']}",
                    key="ref_lookup_idx",
                )
                ref_row = results[ref_idx]
                ref_trade, ref_unwind = results_raw[ref_idx]

                col_ref1, col_ref2 = st.columns(2)
                with col_ref1:
                    st.write(f"Underlying: {ref_row['Underlying (code)']}")
                    if st.button("Last Price 조회 (yfinance)", key="fetch_last_price"):
                        result, err = fetch_reference_last_price(ref_row["Underlying (code)"])
                        if err:
                            st.error(err)
                        else:
                            price, price_date, yahoo_ticker = result
                            st.success(f"{yahoo_ticker} 종가 {price} ({price_date} 기준)")
                            st.caption(
                                "참고: Unwind 계산엔 Last Price가 직접 쓰이지 않아요(청산가는 "
                                "실제 체결가를 그대로 씀). Reset 처리 만들 때 여기서 쓰일 값이에요."
                            )

                with col_ref2:
                    st.write(f"Floating Rate Index: {ref_row.get('Floating Rate Index')}")
                    ecos_key = st.text_input(
                        "ECOS 인증키", type="password", key="ecos_api_key",
                        help="https://ecos.bok.or.kr/api/ 에서 무료 발급",
                    )
                    ref_date = st.date_input(
                        "조회 날짜", value=pd.Timestamp.now().date(), key="ecos_ref_date"
                    )
                    if st.button("CD91 금리 조회 + 참고용 계산 (ECOS)", key="fetch_cd91"):
                        if not ecos_key:
                            st.warning("ECOS 인증키를 먼저 입력해주세요.")
                        else:
                            rate, err = fetch_reference_cd91_rate(
                                ecos_key, ref_date.strftime("%Y%m%d")
                            )
                            if err:
                                st.error(err)
                            else:
                                st.success(f"CD91 금리: {rate}%")

                                ref_unwind_with_rate = dict(ref_unwind)
                                ref_unwind_with_rate["interest_rate_override"] = rate / 100
                                ref_calc = compute_unwind_result(
                                    ref_trade, ref_unwind_with_rate, holidays
                                )
                                st.markdown("**이 금리로 계산하면 (참고용, 실제 값에 반영 안 됨):**")
                                mcol1, mcol2, mcol3 = st.columns(3)
                                mcol1.metric(
                                    "Floating Rate Payments",
                                    f"{ref_calc['Floating Rate Payments']:,}"
                                    if ref_calc["Floating Rate Payments"] is not None else "-",
                                )
                                mcol2.metric(
                                    "Unwind Payment",
                                    f"{ref_calc['Unwind Payment (Swap CCY)']:,}"
                                    if ref_calc["Unwind Payment (Swap CCY)"] is not None else "-",
                                )
                                mcol3.metric("Direction", ref_calc["Direction"] or "-")
                                st.caption(
                                    "실제로 쓰려면 위 UNWIND 표의 'Interest Rate 수동입력' 칸에 "
                                    "이 금리를 직접 입력해서 확정해주세요."
                                )

            col_a, col_b = st.columns(2)
            with col_a:
                unwind_raw_bytes = build_unwind_raw_workbook(results)
                st.download_button(
                    "📥 Unwind_Raw 붙여넣기용 엑셀 (A~Z 전체)",
                    data=unwind_raw_bytes,
                    file_name="unwind_raw_export.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
                st.caption(
                    "기존 Unwind_Raw 시트와 같은 컬럼 순서(A~Z)로 6행 헤더+7행부터 데이터가 "
                    "들어있어요. 그대로 복사해서 Unwind_Raw에 붙여넣고, 매크로는 블룸버그 PC에서 "
                    "직접 돌리시면 됩니다."
                )

            with col_b:
                export_rows = [
                    {
                        "trs_code": r["KIS Ref. No."],
                        "floating_rate_index": r["Floating Rate Index"],
                        "fixing_date": r["Interest Fixing Date"],
                        "spread": r["Spread"],
                    }
                    for r in results
                ]
                xlsx_bytes = build_bloomberg_formula_workbook(export_rows)
                st.download_button(
                    "📥 Interest Rate 블룸버그 수식 엑셀",
                    data=xlsx_bytes,
                    file_name="interest_rate_bloomberg.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
                st.caption(
                    "수식만 있는 파일이에요. 블룸버그 PC에서 열면 값이 계산되고, 그 값을 위 UNWIND "
                    "표의 'Interest Rate 수동입력' 칸에 옮겨 적으면 이자 관련 계산까지 완성돼요."
                )
        else:
            st.info("TRS Code가 거래 표(OPEN)에 있는 KIS Ref. No.와 매칭되는 UNWIND 건이 없습니다.")

elif page == "블룸버그 대체 데이터 확인":
    st.subheader("🧪 블룸버그 대체 데이터 확인 (yfinance / ECOS)")
    st.caption(
        "Reset 처리용 Last Price(yfinance)와 CD91일물 금리(ECOS)를 미리 확인/검증하는 "
        "페이지예요. 아직 Reset 계산에 자동 연동은 안 돼 있고, 값 확인용입니다."
    )

    verify_ssl = not st.checkbox(
        "SSL 인증서 검증 건너뛰기 (사내망 프록시 때문에 SSL 에러 날 때만 체크)",
        value=False,
    )
    if not verify_ssl:
        st.warning(
            "⚠️ 인증서 검증을 끈 상태예요 — 문제 없이 되면 다시 꺼두는 게 안전해요. "
            "가능하면 'pip install pip-system-certs'로 근본 해결을 추천드려요."
        )

    st.markdown("#### 1) 종목 Last Price (yfinance)")
    trades_for_check = load_trades(book=book_filter, counterparty=counterparty_filter)
    if trades_for_check.empty:
        st.info("등록된 거래가 없어요.")
    else:
        pick = st.selectbox(
            "거래 선택",
            trades_for_check["kis_ref_no"].dropna().tolist(),
        )
        row = trades_for_check[trades_for_check["kis_ref_no"] == pick].iloc[0]
        yf_ticker = bloomberg_to_yfinance_ticker(row["underlying_code"])
        st.write(f"Underlying: `{row['underlying_code']}` → yfinance 티커: `{yf_ticker}`")

        price_target_date = st.date_input(
            "기준일 (기본: 어제 — 휴일이면 그 이전 가장 최근 거래일 종가를 자동으로 가져와요)",
            value=pd.Timestamp.now().date() - pd.Timedelta(days=1),
        )

        if yf_ticker and st.button("Last Price 조회", key="fetch_yf"):
            price, date, company_name, err = fetch_yf_price_asof(
                yf_ticker, price_target_date, verify_ssl=verify_ssl
            )
            if err:
                st.error(f"조회 실패: {err}")
            else:
                if date != str(price_target_date):
                    st.caption(f"({price_target_date}은 휴일/주말이라 {date}의 종가로 대체됨)")
                st.success(f"{date} 종가: {price}")
                st.write(f"yfinance 종목명: {company_name}")
                stored_name = (row["underlying_description"] or "").upper()
                if company_name and stored_name and not (
                    any(w in stored_name for w in company_name.upper().split()[:2])
                    or any(w in company_name.upper() for w in stored_name.split()[:2])
                ):
                    st.warning(
                        f"⚠️ 저장된 종목명('{row['underlying_description']}')과 "
                        f"yfinance 종목명('{company_name}')이 많이 달라 보여요 — 티커 확인 필요"
                    )
                else:
                    st.caption("종목명 대략 일치 — 티커가 맞는 것 같아요.")
        elif not yf_ticker:
            st.warning("이 종목은 자동 티커 변환 대상이 아니에요(중국 A주만 지원). 수동 확인 필요.")

    st.divider()
    st.markdown("#### 2) ECOS 통계 항목 코드 확인")
    st.caption(
        "정확한 CD91일물 항목코드를 확실히 하기 위해, 통계표 안의 전체 항목을 조회해서 "
        "직접 눈으로 확인하는 기능이에요. 통계표코드는 보통 시장금리는 '817Y002'입니다."
    )
    stat_code_input = st.text_input("통계표코드", value="817Y002")
    if st.button("항목 목록 조회", key="fetch_ecos_items"):
        try:
            items_df = ecos_list_items(stat_code_input, verify_ssl=verify_ssl)
            if items_df.empty:
                st.warning("항목이 없어요 — 통계표코드를 확인해주세요.")
            else:
                st.dataframe(items_df, use_container_width=True, hide_index=True)
        except Exception as e:
            st.error(f"조회 실패: {e}")

    st.divider()
    st.markdown("#### 3) 특정 항목 값 조회")
    col_a, col_b, col_c = st.columns(3)
    with col_a:
        item_code_input = st.text_input("항목코드", value="010502000")
    with col_b:
        ecos_date = st.date_input("조회일자", value=pd.Timestamp.now().date())
    with col_c:
        st.write("")
        st.write("")
        if st.button("값 조회", key="fetch_ecos_value"):
            value, err = fetch_ecos_series(
                stat_code_input, item_code_input, ecos_date.strftime("%Y%m%d"),
                verify_ssl=verify_ssl,
            )
            if err:
                st.error(err)
            else:
                st.success(f"{ecos_date} 값: {value}")

elif page == "휴일 관리(WORKDAY)":
    st.subheader("휴일 캘린더 (WORKDAY)")
    st.caption(
        "엑셀 WORKDAY 시트를 그대로 붙여넣으세요 — A열(날짜), B열(휴일명) 순서 그대로입니다. "
        "여기 등록된 날짜는 Settlement Date 계산(WORKDAY 함수 대체)에서 자동으로 건너뜁니다."
    )

    holiday_df = load_holiday_editor_df()

    edited_holiday_df = st.data_editor(
        holiday_df,
        num_rows="dynamic",
        use_container_width=True,
        hide_index=True,
        key="holiday_editor",
        column_config={
            "id": st.column_config.NumberColumn("ID", disabled=True),
        },
    )

    h_state = st.session_state.get("holiday_editor", {})
    h_add = len(h_state.get("added_rows", []))
    h_edit = len(h_state.get("edited_rows", {}))
    h_del = len(h_state.get("deleted_rows", []))
    st.caption(f"[추가] {h_add}  [변경] {h_edit}  [삭제] {h_del}")

    if st.button("Save Changes", type="primary", key="save_holiday"):
        n_added, n_edited, n_deleted = apply_holiday_editor_changes(holiday_df, h_state)
        st.success(f"저장 완료 — 추가 {n_added} / 변경 {n_edited} / 삭제 {n_deleted}")
        st.rerun()

else:
    st.subheader("거래 조회")
    col1, col2, col3 = st.columns(3)
    with col1:
        f_counterparty = st.text_input("Counterparty 검색")
    with col2:
        book_options = ["전체", "End Client", "BTB"]
        f_book = st.selectbox("Book", book_options, index=book_options.index(book_filter))
    with col3:
        f_status = st.selectbox("상태", ["전체", "open", "reset", "unwind", "rollover"])

    df = load_trades(f_counterparty, f_book, f_status)
    st.caption(f"총 {len(df)}건")
    if not df.empty:
        unwound = load_unwound_totals()
        df["Balance (Net No of Units)"] = df.apply(
            lambda r: (r["number_of_units"] or 0) - unwound.get(r["kis_ref_no"], 0),
            axis=1,
        )
        df.insert(0, "No.", range(1, len(df) + 1))
    st.dataframe(df, use_container_width=True, hide_index=True)
    st.caption(
        "Balance (Net No of Units) = OPEN(최종본) AF열과 동일한 로직 — 계약 수량(Number of "
        "Units)에서 UNWIND 입력/수정에 기록된 청산 수량 누적분을 뺀 값이에요."
    )
